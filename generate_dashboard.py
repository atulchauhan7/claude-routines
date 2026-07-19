#!/usr/bin/env python3
"""
Build 'Daily Store & Ads Performance Sheet.xlsx' from the JSON data stores in data/.

Historical daily rows (data/shopify_daily_history.json, data/meta_daily_history.json)
are append-only, deduped by date, and written verbatim into the two "Daily Data" tabs.
The Dashboard tab reads from those tabs with formulas (INDEX/OFFSET/AVERAGE), so it stays
live as new days are appended by future runs.

The Product Performance / Campaign Performance / Recommendations tabs are populated from
data/snapshot_latest.json, which holds the most recent day's product & campaign/ad-set
breakdown and the generated recommendations.

Usage: python3 generate_dashboard.py
Output: output/Daily_Store_and_Ads_Performance_Sheet.xlsx
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "data")
OUT = os.path.join(BASE, "output")
os.makedirs(OUT, exist_ok=True)

FONT_NAME = "Arial"
CURRENCY_FMT = '"₹"#,##0.00'
CURRENCY_FMT0 = '"₹"#,##0'
PCT_FMT = '0.00"%"'
ROAS_FMT = '0.00"x"'

C_HEADER_BG = "1F1F1F"
C_HEADER_FG = "FFFFFF"
C_ACCENT = "2A6FF2"
C_ACCENT_FG = "FFFFFF"
C_SECTION_BG = "F2F2F2"
C_GREEN = "C6E8C6"
C_RED = "F5B6B6"
C_YELLOW = "FFF3CC"

HEADER_FONT = Font(name=FONT_NAME, bold=True, color=C_HEADER_FG, size=11)
HEADER_FILL = PatternFill("solid", fgColor=C_HEADER_BG)
ACCENT_FONT = Font(name=FONT_NAME, bold=True, color=C_ACCENT_FG, size=13)
ACCENT_FILL = PatternFill("solid", fgColor=C_ACCENT)
SECTION_FONT = Font(name=FONT_NAME, bold=True, size=11)
SECTION_FILL = PatternFill("solid", fgColor=C_SECTION_BG)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=16, color=C_ACCENT_FG)
NORMAL_FONT = Font(name=FONT_NAME, size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_json(name):
    with open(os.path.join(DATA, name)) as f:
        return json.load(f)


def style_header_row(ws, row, ncols, fill=HEADER_FILL, font=HEADER_FONT):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def autofit(ws, ncols, min_width=10, max_width=48):
    for c in range(1, ncols + 1):
        letter = get_column_letter(c)
        longest = 0
        for cell in ws[letter]:
            if cell.value is not None:
                longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, longest + 3))


def apply_all_borders(ws, min_row, max_row, min_col, max_col):
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(row=r, column=c).border = BORDER


# ─────────────────────────── Shopify Daily Data ───────────────────────────

def build_shopify_daily(wb, shopify_hist):
    ws = wb.create_sheet("Shopify Daily Data")
    headers = ["Date", "Gross Sales", "Net Sales", "Orders", "AOV", "Total Sales",
               "Discounts", "Returns", "Shipping", "Taxes", "Sessions", "Cart Adds",
               "Reached Checkout", "Completed Checkout", "Conversion Rate"]
    ws.append(headers)
    money_cols = {2, 3, 5, 6, 7, 8, 9, 10}
    pct_cols = {15}
    for row in shopify_hist:
        ws.append([
            row["date"], row["gross_sales"], row["net_sales"], row["orders"], row["aov"],
            row["total_sales"], row["discounts"], row["returns"], row["shipping"],
            row["taxes"], row["sessions"], row["cart_adds"], row["reached_checkout"],
            row["completed_checkout"],
            (row["conversion_rate"] / 100.0) if row["conversion_rate"] is not None else None,
        ])
    ncols = len(headers)
    nrows = ws.max_row
    style_header_row(ws, 1, ncols)
    for r in range(2, nrows + 1):
        for c in money_cols:
            ws.cell(row=r, column=c).number_format = CURRENCY_FMT
        for c in pct_cols:
            ws.cell(row=r, column=c).number_format = "0.00%"
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, 1, nrows, 1, ncols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"
    autofit(ws, ncols)

    # Conditional formatting: highlight big single-day return spikes (Returns column, H)
    ws.conditional_formatting.add(
        f"H2:H{nrows}",
        CellIsRule(operator="lessThan", formula=["-20000"], fill=PatternFill("solid", fgColor=C_RED)),
    )
    # Orders low vs typical band
    ws.conditional_formatting.add(
        f"D2:D{nrows}",
        ColorScaleRule(start_type="min", start_color="F5B6B6",
                        mid_type="percentile", mid_value=50, mid_color="FFF3CC",
                        end_type="max", end_color="C6E8C6"),
    )
    return ws, nrows, ncols


# ─────────────────────────── Meta Ads Daily Data ───────────────────────────

def build_meta_daily(wb, meta_hist):
    ws = wb.create_sheet("Meta Ads Daily Data")
    headers = ["Date", "Spend", "Impressions", "Reach", "Clicks", "CTR", "CPC", "CPM",
               "Purchases", "CPA", "ROAS"]
    ws.append(headers)
    money_cols = {2, 7, 8, 10}
    pct_cols = {6}
    roas_cols = {11}
    for row in meta_hist:
        ws.append([
            row["date"], row["spend"], row["impressions"], row["reach"], row["clicks"],
            row["ctr"] / 100.0, row["cpc"], row["cpm"], row["purchases"], row["cpa"], row["roas"],
        ])
    ncols = len(headers)
    nrows = ws.max_row
    style_header_row(ws, 1, ncols)
    for r in range(2, nrows + 1):
        for c in money_cols:
            ws.cell(row=r, column=c).number_format = CURRENCY_FMT
        for c in pct_cols:
            ws.cell(row=r, column=c).number_format = "0.00%"
        for c in roas_cols:
            ws.cell(row=r, column=c).number_format = ROAS_FMT
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, 1, nrows, 1, ncols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"
    autofit(ws, ncols)

    # Conditional formatting: ROAS band (col K)
    ws.conditional_formatting.add(
        f"K2:K{nrows}",
        CellIsRule(operator="lessThan", formula=["1.5"], fill=PatternFill("solid", fgColor=C_RED)),
    )
    ws.conditional_formatting.add(
        f"K2:K{nrows}",
        CellIsRule(operator="between", formula=["1.5", "3"], fill=PatternFill("solid", fgColor=C_YELLOW)),
    )
    ws.conditional_formatting.add(
        f"K2:K{nrows}",
        CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid", fgColor=C_GREEN)),
    )
    return ws, nrows, ncols


# ─────────────────────────── Product Performance ───────────────────────────

def build_product_performance(wb, snap):
    ws = wb.create_sheet("Product Performance")
    ws["A1"] = f"PRODUCT PERFORMANCE — Yesterday ({snap['date']}) vs Trailing 7 Days"
    ws["A1"].font = ACCENT_FONT
    ws.merge_cells("A1:E1")
    for c in range(1, 6):
        ws.cell(row=1, column=c).fill = ACCENT_FILL

    headers1 = ["Product", "Gross Sales", "Net Sales", "Orders", "Total Sales"]
    ws.append([])
    ws.append(headers1)
    hdr_row1 = ws.max_row
    for row in snap["products_yesterday"]:
        ws.append(row)
    end1 = ws.max_row
    style_header_row(ws, hdr_row1, 5)
    for r in range(hdr_row1 + 1, end1 + 1):
        for c in (2, 3, 5):
            ws.cell(row=r, column=c).number_format = CURRENCY_FMT
        for c in range(1, 6):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, hdr_row1, end1, 1, 5)
    ws.auto_filter.ref = f"A{hdr_row1}:E{end1}"

    ws.append([])
    ws.append(["7-DAY TOP PRODUCTS (2026-07-11 to 2026-07-17)"])
    title2_row = ws.max_row
    ws.cell(row=title2_row, column=1).font = SECTION_FONT
    ws.cell(row=title2_row, column=1).fill = SECTION_FILL
    ws.merge_cells(f"A{title2_row}:E{title2_row}")
    for c in range(1, 6):
        ws.cell(row=title2_row, column=c).fill = SECTION_FILL

    ws.append(["Product", "Gross Sales", "Net Sales", "Orders", "Total Sales"])
    hdr_row2 = ws.max_row
    for row in snap["products_7day"]:
        ws.append(row)
    end2 = ws.max_row
    style_header_row(ws, hdr_row2, 5)
    for r in range(hdr_row2 + 1, end2 + 1):
        for c in (2, 3, 5):
            ws.cell(row=r, column=c).number_format = CURRENCY_FMT
        for c in range(1, 6):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, hdr_row2, end2, 1, 5)

    ws.freeze_panes = "A3"
    autofit(ws, 5, min_width=14, max_width=60)
    return ws, hdr_row1, end1


# ─────────────────────────── Campaign Performance ───────────────────────────

def build_campaign_performance(wb, snap):
    ws = wb.create_sheet("Campaign Performance")
    ws["A1"] = f"META ADS — CAMPAIGN & AD SET PERFORMANCE ({snap['date']})"
    ws["A1"].font = ACCENT_FONT
    ncols = 12
    ws.merge_cells(f"A1:{get_column_letter(ncols)}1")
    for c in range(1, ncols + 1):
        ws.cell(row=1, column=c).fill = ACCENT_FILL

    ws.append([])
    ws.append(["Campaign", "Status", "Spend", "Impressions", "Reach", "Clicks",
               "CTR", "CPC", "CPM", "Purchases", "CPA", "ROAS"])
    camp_hdr = ws.max_row
    camps = sorted(snap["campaigns_yesterday"], key=lambda c: c["spend"], reverse=True)
    for c in camps:
        ws.append([c["name"], c["status"], c["spend"], c["impressions"], c["reach"],
                   c["clicks"], c["ctr"] / 100.0, c["cpc"], c["cpm"], c["purchases"],
                   c["cpa"], c["roas"]])
    camp_end = ws.max_row
    style_header_row(ws, camp_hdr, ncols)
    for r in range(camp_hdr + 1, camp_end + 1):
        ws.cell(row=r, column=3).number_format = CURRENCY_FMT
        ws.cell(row=r, column=7).number_format = "0.00%"
        ws.cell(row=r, column=8).number_format = CURRENCY_FMT
        ws.cell(row=r, column=9).number_format = CURRENCY_FMT
        ws.cell(row=r, column=11).number_format = CURRENCY_FMT
        ws.cell(row=r, column=12).number_format = ROAS_FMT
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, camp_hdr, camp_end, 1, ncols)
    ws.auto_filter.ref = f"A{camp_hdr}:{get_column_letter(ncols)}{camp_end}"
    # ROAS conditional format for campaigns
    roas_col = get_column_letter(12)
    ws.conditional_formatting.add(f"{roas_col}{camp_hdr+1}:{roas_col}{camp_end}",
        CellIsRule(operator="lessThan", formula=["1.5"], fill=PatternFill("solid", fgColor=C_RED)))
    ws.conditional_formatting.add(f"{roas_col}{camp_hdr+1}:{roas_col}{camp_end}",
        CellIsRule(operator="between", formula=["1.5", "3"], fill=PatternFill("solid", fgColor=C_YELLOW)))
    ws.conditional_formatting.add(f"{roas_col}{camp_hdr+1}:{roas_col}{camp_end}",
        CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid", fgColor=C_GREEN)))
    # CTR underperform (<1.5%) flag
    ctr_col = get_column_letter(7)
    ws.conditional_formatting.add(f"{ctr_col}{camp_hdr+1}:{ctr_col}{camp_end}",
        CellIsRule(operator="lessThan", formula=["0.015"], fill=PatternFill("solid", fgColor=C_YELLOW)))

    ws.append([])
    ws.append(["AD SET LEVEL PERFORMANCE"])
    as_title = ws.max_row
    ws.merge_cells(f"A{as_title}:{get_column_letter(ncols)}{as_title}")
    for c in range(1, ncols + 1):
        ws.cell(row=as_title, column=c).fill = SECTION_FILL
    ws.cell(row=as_title, column=1).font = SECTION_FONT

    ws.append(["Ad Set", "Campaign", "Spend", "Impressions", "Reach", "Clicks",
               "CTR", "CPC", "CPM", "Purchases", "CPA", "ROAS"])
    as_hdr = ws.max_row
    adsets = sorted(snap["adsets_yesterday"], key=lambda a: a["spend"], reverse=True)
    for a in adsets:
        ws.append([a["name"], a["campaign"], a["spend"], a["impressions"], a["reach"],
                   a["clicks"], a["ctr"] / 100.0, a["cpc"], a["cpm"], a["purchases"],
                   a["cpa"], a["roas"]])
    as_end = ws.max_row
    style_header_row(ws, as_hdr, ncols)
    for r in range(as_hdr + 1, as_end + 1):
        ws.cell(row=r, column=3).number_format = CURRENCY_FMT
        ws.cell(row=r, column=7).number_format = "0.00%"
        ws.cell(row=r, column=8).number_format = CURRENCY_FMT
        ws.cell(row=r, column=9).number_format = CURRENCY_FMT
        ws.cell(row=r, column=11).number_format = CURRENCY_FMT
        ws.cell(row=r, column=12).number_format = ROAS_FMT
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).font = NORMAL_FONT
    apply_all_borders(ws, as_hdr, as_end, 1, ncols)
    ws.auto_filter.ref = f"A{as_hdr}:{get_column_letter(ncols)}{as_end}"
    ws.conditional_formatting.add(f"{roas_col}{as_hdr+1}:{roas_col}{as_end}",
        CellIsRule(operator="lessThan", formula=["1.5"], fill=PatternFill("solid", fgColor=C_RED)))
    ws.conditional_formatting.add(f"{roas_col}{as_hdr+1}:{roas_col}{as_end}",
        CellIsRule(operator="between", formula=["1.5", "3"], fill=PatternFill("solid", fgColor=C_YELLOW)))
    ws.conditional_formatting.add(f"{roas_col}{as_hdr+1}:{roas_col}{as_end}",
        CellIsRule(operator="greaterThan", formula=["3"], fill=PatternFill("solid", fgColor=C_GREEN)))

    ws.freeze_panes = f"A{camp_hdr + 1}"
    autofit(ws, ncols, min_width=10, max_width=34)
    return ws


# ─────────────────────────── Recommendations & Notes ───────────────────────────

def build_recommendations(wb, snap, recs):
    ws = wb.create_sheet("Recommendations & Notes")
    ws["A1"] = f"RECOMMENDATIONS & NOTES — {snap['date']}"
    ws["A1"].font = ACCENT_FONT
    ws.merge_cells("A1:C1")
    for c in range(1, 4):
        ws.cell(row=1, column=c).fill = ACCENT_FILL
    ws.append([])

    def section(title, items, tag_colors=None):
        ws.append([title])
        r = ws.max_row
        ws.cell(row=r, column=1).font = SECTION_FONT
        ws.cell(row=r, column=1).fill = SECTION_FILL
        ws.merge_cells(f"A{r}:C{r}")
        for c in range(1, 4):
            ws.cell(row=r, column=c).fill = SECTION_FILL
        for item in items:
            ws.append(["", item, ""])
            rr = ws.max_row
            ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=rr, column=2).font = NORMAL_FONT
        ws.append([])

    section("KEY WINS", recs["wins"])
    section("KEY ISSUES", recs["issues"])
    section("RECOMMENDED ACTIONS FOR TODAY", [f"{i+1}. {a}" for i, a in enumerate(recs["actions"])])
    section("ANOMALY SIGNALS (Meta Ads)", snap.get("anomalies", []) or ["None detected."])
    section("TRACKING NOTE", [snap.get("tracking_issue_note", "")])
    section("UNAVAILABLE METRICS / DATA NOTES", snap.get("unavailable_metrics", []))

    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 100
    ws.column_dimensions["C"].width = 4
    ws.freeze_panes = "A3"
    return ws


# ─────────────────────────── Dashboard ───────────────────────────

def build_dashboard(wb, shopify_nrows, shopify_ncols, meta_nrows, meta_ncols, snap, recs):
    ws = wb.create_sheet("Dashboard", 0)
    ws["A1"] = "DAILY STORE & ADS PERFORMANCE DASHBOARD"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")
    for c in range(1, 7):
        ws.cell(row=1, column=c).fill = ACCENT_FILL
    ws["A2"] = f"Reporting date (Asia/Kolkata): {snap['date']}   |   Store: {snap['shop_domain']}   |   Currency: {snap['currency']}"
    ws["A2"].font = Font(name=FONT_NAME, italic=True, size=10)
    ws.merge_cells("A2:F2")

    SD = "'Shopify Daily Data'"
    MD = "'Meta Ads Daily Data'"

    def last_row_formula(sheet):
        return f"COUNTA({sheet}!A:A)"

    row = 4
    ws.cell(row=row, column=1, value="EXECUTIVE SUMMARY").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(f"A{row}:F{row}")
    for c in range(1, 7):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1
    summary_lines = [
        "Both Shopify revenue and Meta Ads purchases fell sharply yesterday vs. the trailing 7-day average.",
        "Ad efficiency (CPC, CPM) improved, but ROAS and CPA moved the wrong way as purchase volume dropped.",
        "One ad set is running below breakeven and one is flagged for a narrow audience — see Campaign Performance tab.",
    ]
    for line in summary_lines:
        ws.cell(row=row, column=1, value=line).font = NORMAL_FONT
        ws.merge_cells(f"A{row}:F{row}")
        row += 1
    row += 1

    # ── Shopify performance table (formulas) ──
    ws.cell(row=row, column=1, value="SHOPIFY PERFORMANCE").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(f"A{row}:D{row}")
    for c in range(1, 5):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1
    hdr_row = row
    ws.append_row = None  # no-op guard
    for i, h in enumerate(["Metric", "Yesterday", "7-Day Avg", "Change vs Avg"]):
        ws.cell(row=hdr_row, column=1 + i, value=h)
    style_header_row(ws, hdr_row, 4)
    row += 1

    shopify_metrics = [
        ("Gross Sales", "B", CURRENCY_FMT),
        ("Net Sales", "C", CURRENCY_FMT),
        ("Orders", "D", "#,##0"),
        ("Avg Order Value", "E", CURRENCY_FMT),
        ("Returns", "H", CURRENCY_FMT),
    ]
    sd_first_row = 2
    lastN = "SDLAST"
    for label, col, fmt in shopify_metrics:
        r = row
        ws.cell(row=r, column=1, value=label).font = NORMAL_FONT
        yest_formula = f"=INDEX({SD}!{col}:{col},{last_row_formula(SD)})"
        avg_formula = (f"=IFERROR(AVERAGE(OFFSET({SD}!{col}1,{last_row_formula(SD)}-8,0,7,1)),"
                        f"\"n/a\")")
        chg_formula = f"=IFERROR((B{r}-C{r})/C{r},\"n/a\")"
        ws.cell(row=r, column=2, value=yest_formula).number_format = fmt
        ws.cell(row=r, column=3, value=avg_formula).number_format = fmt
        ws.cell(row=r, column=4, value=chg_formula).number_format = "+0.0%;-0.0%"
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).font = NORMAL_FONT
        row += 1
    shopify_block_end = row - 1
    row += 1

    # ── Meta Ads performance table (formulas) ──
    ws.cell(row=row, column=1, value="META ADS PERFORMANCE").font = SECTION_FONT
    ws.cell(row=row, column=1).fill = SECTION_FILL
    ws.merge_cells(f"A{row}:D{row}")
    for c in range(1, 5):
        ws.cell(row=row, column=c).fill = SECTION_FILL
    row += 1
    hdr_row2 = row
    for i, h in enumerate(["Metric", "Yesterday", "7-Day Avg", "Change vs Avg"]):
        ws.cell(row=hdr_row2, column=1 + i, value=h)
    style_header_row(ws, hdr_row2, 4)
    row += 1

    meta_metrics = [
        ("Spend", "B", CURRENCY_FMT),
        ("Purchases", "I", "#,##0"),
        ("CPA", "J", CURRENCY_FMT),
        ("ROAS", "K", ROAS_FMT),
        ("CTR", "F", "0.00%"),
        ("CPC", "G", CURRENCY_FMT),
    ]
    for label, col, fmt in meta_metrics:
        r = row
        ws.cell(row=r, column=1, value=label).font = NORMAL_FONT
        yest_formula = f"=INDEX({MD}!{col}:{col},{last_row_formula(MD)})"
        avg_formula = (f"=IFERROR(AVERAGE(OFFSET({MD}!{col}1,{last_row_formula(MD)}-8,0,7,1)),"
                        f"\"n/a\")")
        chg_formula = f"=IFERROR((B{r}-C{r})/C{r},\"n/a\")"
        ws.cell(row=r, column=2, value=yest_formula).number_format = fmt
        ws.cell(row=r, column=3, value=avg_formula).number_format = fmt
        ws.cell(row=r, column=4, value=chg_formula).number_format = "+0.0%;-0.0%"
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).font = NORMAL_FONT
        row += 1
    meta_block_end = row - 1
    row += 1

    # Conditional formatting: change columns green/red
    ws.conditional_formatting.add(f"D{hdr_row+1}:D{shopify_block_end}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=C_GREEN)))
    ws.conditional_formatting.add(f"D{hdr_row+1}:D{shopify_block_end}",
        CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=C_RED)))
    ws.conditional_formatting.add(f"D{hdr_row2+1}:D{meta_block_end}",
        CellIsRule(operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor=C_GREEN)))
    ws.conditional_formatting.add(f"D{hdr_row2+1}:D{meta_block_end}",
        CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=C_RED)))

    # ── Key wins / issues / actions ──
    def bullet_block(title, items):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = SECTION_FONT
        ws.cell(row=row, column=1).fill = SECTION_FILL
        ws.merge_cells(f"A{row}:F{row}")
        for c in range(1, 7):
            ws.cell(row=row, column=c).fill = SECTION_FILL
        row += 1
        for item in items:
            ws.cell(row=row, column=1, value="•")
            ws.cell(row=row, column=2, value=item).font = NORMAL_FONT
            ws.merge_cells(f"B{row}:F{row}")
            ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True)
            row += 1
        row += 1

    bullet_block("KEY WINS", recs["wins"])
    bullet_block("KEY ISSUES", recs["issues"])
    bullet_block("RECOMMENDED ACTIONS FOR TODAY", [f"{i+1}. {a}" for i, a in enumerate(recs["actions"])])

    chart_anchor_row = row + 1

    ws.freeze_panes = "A4"
    ws.column_dimensions["A"].width = 26
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20

    # ── Charts ──
    # Sales trend (Gross Sales) + Orders trend
    sales_chart = LineChart()
    sales_chart.title = "Gross Sales Trend (₹)"
    sales_chart.y_axis.title = "Gross Sales"
    sales_chart.x_axis.title = "Date"
    sales_chart.height, sales_chart.width = 8, 15
    data_ref = Reference(wb["Shopify Daily Data"], min_col=2, min_row=1, max_row=shopify_nrows)
    cats_ref = Reference(wb["Shopify Daily Data"], min_col=1, min_row=2, max_row=shopify_nrows)
    sales_chart.add_data(data_ref, titles_from_data=True)
    sales_chart.set_categories(cats_ref)
    ws.add_chart(sales_chart, f"H4")

    orders_chart = BarChart()
    orders_chart.type = "col"
    orders_chart.title = "Orders Trend"
    orders_chart.height, orders_chart.width = 8, 15
    data_ref = Reference(wb["Shopify Daily Data"], min_col=4, min_row=1, max_row=shopify_nrows)
    orders_chart.add_data(data_ref, titles_from_data=True)
    orders_chart.set_categories(cats_ref)
    ws.add_chart(orders_chart, "H21")

    spend_chart = LineChart()
    spend_chart.title = "Ad Spend Trend (₹)"
    spend_chart.height, spend_chart.width = 8, 15
    mdata_ref = Reference(wb["Meta Ads Daily Data"], min_col=2, min_row=1, max_row=meta_nrows)
    mcats_ref = Reference(wb["Meta Ads Daily Data"], min_col=1, min_row=2, max_row=meta_nrows)
    spend_chart.add_data(mdata_ref, titles_from_data=True)
    spend_chart.set_categories(mcats_ref)
    ws.add_chart(spend_chart, "Q4")

    roas_chart = LineChart()
    roas_chart.title = "ROAS Trend"
    roas_chart.height, roas_chart.width = 8, 15
    rdata_ref = Reference(wb["Meta Ads Daily Data"], min_col=11, min_row=1, max_row=meta_nrows)
    roas_chart.add_data(rdata_ref, titles_from_data=True)
    roas_chart.set_categories(mcats_ref)
    ws.add_chart(roas_chart, "Q21")

    # Top products bar chart (yesterday) — pull from Product Performance sheet
    pp = wb["Product Performance"]
    top_products_chart = BarChart()
    top_products_chart.type = "bar"
    top_products_chart.title = "Top Products — Yesterday (Gross Sales ₹)"
    top_products_chart.height, top_products_chart.width = 10, 15
    n_products = len(snap["products_yesterday"])
    pdata_ref = Reference(pp, min_col=2, min_row=3, max_row=3 + n_products)
    pcats_ref = Reference(pp, min_col=1, min_row=4, max_row=3 + n_products)
    top_products_chart.add_data(pdata_ref, titles_from_data=True)
    top_products_chart.set_categories(pcats_ref)
    ws.add_chart(top_products_chart, "H38")

    # Top campaigns bar chart
    cp = wb["Campaign Performance"]
    n_camps = len(snap["campaigns_yesterday"])
    top_campaigns_chart = BarChart()
    top_campaigns_chart.type = "bar"
    top_campaigns_chart.title = "Campaign Spend — Yesterday (₹)"
    top_campaigns_chart.height, top_campaigns_chart.width = 10, 15
    cdata_ref = Reference(cp, min_col=3, min_row=3, max_row=3 + n_camps)
    ccats_ref = Reference(cp, min_col=1, min_row=4, max_row=3 + n_camps)
    top_campaigns_chart.add_data(cdata_ref, titles_from_data=True)
    top_campaigns_chart.set_categories(ccats_ref)
    ws.add_chart(top_campaigns_chart, "Q38")

    return ws


def main():
    shopify_hist = load_json("shopify_daily_history.json")
    meta_hist = load_json("meta_daily_history.json")
    snap = load_json("snapshot_latest.json")

    # Dedup by date, sort chronologically ascending (append-only guarantee)
    def dedup_sort(rows):
        seen = {}
        for r in rows:
            seen[r["date"]] = r
        return [seen[d] for d in sorted(seen)]

    shopify_hist = dedup_sort(shopify_hist)
    meta_hist = dedup_sort(meta_hist)

    recs = {
        "wins": [
            "Campaign 'Retargeting- new' delivered ROAS 4.42x yesterday (its ad set 'Best performing ads' hit 5.51x) — a strong scale candidate.",
            "Ad efficiency improved: CPC down 17.2% and CPM down 12.2% vs the 7-day average.",
            "'Ira Rayon Co-ord Set' remains the top seller, ₹29,988 yesterday and ₹307,377 over the trailing 7 days.",
        ],
        "issues": [
            "Gross sales fell 39.3% and orders fell 37.7% vs the 7-day average — the steepest single-day drop of the week.",
            "Meta purchases fell 41.0% and ROAS fell from 3.84x to 2.82x (-26.6%) while CPA rose 33.9% to ₹909.",
            "Ad set 'White Dhurandhar' is running at ROAS 0.95x — below breakeven.",
            "Ad set 'Black and multiple' was flagged by Meta for a narrow audience (~23K-27K reach).",
            "Shopify's checkout-completion tracking shows 0-1 sessions/day all week despite 41-87 real orders/day — looks like a tracking/pixel gap, not a real conversion problem.",
        ],
        "actions": [
            "Scale budget 20-30% on the 'Retargeting- new' campaign / 'Best performing ads' ad set given its 5.51x ROAS.",
            "Pause or refresh creative on the 'White Dhurandhar' ad set (ROAS 0.95x, below breakeven).",
            "Expand the audience (or test Advantage+ audience) on 'Black and multiple' to fix the narrow-audience warning.",
            "Investigate yesterday's steep order/revenue drop — check payment gateway status, site uptime, and ad delivery pacing together.",
            "Review Shopify's checkout/pixel tracking configuration — completed-checkout sessions do not reconcile with real order volume.",
        ],
    }

    wb = Workbook()
    wb.remove(wb.active)

    sd_ws, sd_nrows, sd_ncols = build_shopify_daily(wb, shopify_hist)
    md_ws, md_nrows, md_ncols = build_meta_daily(wb, meta_hist)
    build_product_performance(wb, snap)
    build_campaign_performance(wb, snap)
    build_recommendations(wb, snap, recs)
    build_dashboard(wb, sd_nrows, sd_ncols, md_nrows, md_ncols, snap, recs)

    # Tab order per spec
    order = ["Dashboard", "Shopify Daily Data", "Meta Ads Daily Data",
             "Product Performance", "Campaign Performance", "Recommendations & Notes"]
    wb._sheets = [wb[name] for name in order]
    for name in order:
        wb[name].sheet_view.showGridLines = False

    out_path = os.path.join(OUT, "Daily_Store_and_Ads_Performance_Sheet.xlsx")
    wb.save(out_path)
    print(f"Saved: {out_path}")
    print(f"Shopify Daily Data rows: {sd_nrows-1}, Meta Ads Daily Data rows: {md_nrows-1}")


if __name__ == "__main__":
    main()
